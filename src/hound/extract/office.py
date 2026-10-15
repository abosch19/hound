"""Word, Excel, PowerPoint (OOXML), Excel 97-2003 y OpenDocument."""

from __future__ import annotations

import datetime as dt
import zipfile
from xml.etree import ElementTree as ET

from . import Block, ExtractionError, Extracted

MAX_ROWS_PER_SHEET = 20000


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y") if not isinstance(v, dt.datetime) or not (v.hour or v.minute) else v.strftime("%d/%m/%Y %H:%M")
    return str(v).strip()


def _row_text(values) -> str:
    return " | ".join(c for c in (_cell(v) for v in values) if c)


def extract_docx(path: str) -> Extracted:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = docx.Document(path)
    except Exception as exc:  # noqa: BLE001 — python-docx lanza tipos variados
        raise ExtractionError(f"Word ilegible: {exc}") from exc
    out = Extracted(title=doc.core_properties.title or "")
    n = 0
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                n += 1
                out.blocks.append(Block(text, unit="párrafo", n=n))
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                seen: list[str] = []
                for cell in row.cells:  # las celdas combinadas se repiten
                    t = cell.text.strip()
                    if t and (not seen or seen[-1] != t):
                        seen.append(t)
                if seen:
                    n += 1
                    out.blocks.append(Block(" | ".join(seen), unit="párrafo", n=n))
    extra = []
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                extra.extend(p.text.strip() for p in part.paragraphs if p.text.strip())
            except Exception:  # noqa: BLE001
                pass
    if extra:
        out.blocks.append(Block("\n".join(dict.fromkeys(extra)), group="encabezado y pie"))
    return out


def extract_xlsx(path: str) -> Extracted:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Excel ilegible: {exc}") from exc
    out = Extracted(title=(wb.properties.title or "") if wb.properties else "")
    try:
        for ws in wb.worksheets:
            _sheet_blocks(out, ws.title, ws.iter_rows(values_only=True))
    finally:
        wb.close()
    return out


def extract_xls(path: str) -> Extracted:
    import xlrd

    try:
        wb = xlrd.open_workbook(path, on_demand=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Excel ilegible: {exc}") from exc
    out = Extracted()
    try:
        for sh in wb.sheets():
            rows = (sh.row_values(r) for r in range(sh.nrows))
            _sheet_blocks(out, sh.name, rows)
    finally:
        wb.release_resources()
    return out


def _sheet_blocks(out: Extracted, title: str, rows) -> None:
    header = ""
    group = f"hoja {title}"
    count = 0
    for i, row in enumerate(rows, start=1):
        if count >= MAX_ROWS_PER_SHEET:
            out.notes.append(f"Hoja {title}: solo se indexan las primeras {MAX_ROWS_PER_SHEET} filas con datos")
            break
        text = _row_text(row)
        if not text:
            continue
        count += 1
        if not header:
            header = text
        out.blocks.append(Block(text, group=group, unit="fila", n=i, context=header if count > 1 else ""))


def extract_pptx(path: str) -> Extracted:
    from pptx import Presentation

    try:
        prs = Presentation(path)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"PowerPoint ilegible: {exc}") from exc
    out = Extracted(title=prs.core_properties.title or "")
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            _shape_text(shape, parts)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("Notas: " + notes)
        if parts:
            out.blocks.append(Block("\n".join(parts), group=f"diapositiva {i}"))
    return out


def _shape_text(shape, parts: list[str]) -> None:
    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        parts.append(shape.text_frame.text.strip())
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            t = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if t:
                parts.append(t)
    if shape.shape_type == 6:  # grupo
        for sub in shape.shapes:
            _shape_text(sub, parts)


_ODF = {
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
}


def extract_odf(path: str) -> Extracted:
    try:
        with zipfile.ZipFile(path) as zf:
            root = ET.fromstring(zf.read("content.xml"))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise ExtractionError(f"OpenDocument ilegible: {exc}") from exc
    out = Extracted()
    t, tb, dr = (f"{{{_ODF[k]}}}" for k in ("text", "table", "draw"))
    pages = list(root.iter(dr + "page"))
    tables = list(root.iter(tb + "table"))
    if pages:
        for i, page in enumerate(pages, start=1):
            text = "\n".join("".join(p.itertext()).strip() for p in page.iter(t + "p") if "".join(p.itertext()).strip())
            if text:
                out.blocks.append(Block(text, group=f"diapositiva {i}"))
    elif path.lower().endswith(".ods"):
        for table in tables:
            rows = ([("".join(c.itertext())) for c in r.iter(tb + "table-cell")] for r in table.iter(tb + "table-row"))
            _sheet_blocks(out, table.get(tb + "name", "?"), rows)
    else:
        n = 0
        for el in root.iter():
            if el.tag in (t + "p", t + "h"):
                text = "".join(el.itertext()).strip()
                if text:
                    n += 1
                    out.blocks.append(Block(text, unit="párrafo", n=n))
    return out
