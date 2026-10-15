from __future__ import annotations

import os
import shutil
from email.message import EmailMessage
from pathlib import Path

import pytest

from hound.config import Config
from hound.security import SecurityProvider

HAS_TESSERACT = shutil.which("tesseract") is not None
needs_tesseract = pytest.mark.skipif(not HAS_TESSERACT, reason="Tesseract no instalado")


class FakeSecurity(SecurityProvider):
    """Simula NTFS: reglas por carpeta (la más específica gana), usuarios con grupos y 'Denegar'.

    El "descriptor" es texto: b"allow:SID,SID;deny:SID".
    """

    def __init__(self, root: str):
        self.root = root
        self.rules: dict[str, tuple[set[str], set[str]] | None] = {"": ({"S-EVERYONE"}, set())}
        self.groups: dict[str, set[str]] = {}
        self.accounts: dict[str, tuple[str, str]] = {}
        self.share: bytes | None = None
        self.broken_sids: set[str] = set()

    def add_user(self, name: str, sid: str, groups: set[str] = frozenset()):
        self.accounts[name.lower()] = (f"EMPRESA\\{name}", sid)
        self.groups[sid] = {"S-EVERYONE", *groups}

    def set_rule(self, folder: str, allow: set[str] | None, deny: set[str] = frozenset()):
        self.rules[folder] = None if allow is None else (set(allow), set(deny))

    def read_sd(self, path: str) -> bytes | None:
        rel = os.path.relpath(path, self.root).replace("\\", "/")
        best = max((k for k in self.rules if k == "" or rel == k or rel.startswith(k + "/")), key=len)
        rule = self.rules[best]
        if rule is None:
            return None
        allow, deny = rule
        return f"allow:{','.join(sorted(allow))};deny:{','.join(sorted(deny))}".encode()

    def check(self, sid: str, sd: bytes) -> bool:
        if sid in self.broken_sids:
            raise RuntimeError("fallo simulado de Authz")
        allow_s, deny_s = sd.decode().split(";")
        allow = set(filter(None, allow_s[6:].split(",")))
        deny = set(filter(None, deny_s[5:].split(",")))
        ids = {sid, *self.groups.get(sid, set())}
        return bool(ids & allow) and not (ids & deny)

    def lookup_account(self, name: str) -> tuple[str, str]:
        key = name.lower().split("\\")[-1]
        if key not in self.accounts:
            raise RuntimeError(f"No existe la cuenta '{name}'")
        return self.accounts[key]

    def share_sd(self, share_name: str) -> bytes | None:
        return self.share


def make_pdf(path: Path, pages: list[str]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(60, y, line)
            y -= 16
        c.showPage()
    c.save()


def text_image(text: str, size=(1400, 500)):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("L", size, 255)
    d = ImageDraw.Draw(img)
    font = None
    for cand in ("/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
                 "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if os.path.exists(cand):
            font = ImageFont.truetype(cand, 44)
            break
    font = font or ImageFont.load_default(size=44)
    y = 40
    for line in text.split("\n"):
        d.text((40, y), line, fill=0, font=font)
        y += 70
    return img


def build_corpus(root: Path) -> None:
    import docx
    import ezdxf
    import openpyxl
    from pptx import Presentation

    (root / "Proveedores" / "Limpieza").mkdir(parents=True)
    (root / "Contabilidad" / "Facturas").mkdir(parents=True)
    (root / "Obras" / "Reforma oficina").mkdir(parents=True)
    (root / "RRHH" / "Nominas").mkdir(parents=True)
    (root / "Correo").mkdir()

    make_pdf(root / "Proveedores" / "Limpieza" / "contrato_limpieza_2024.pdf", [
        "CONTRATO DE PRESTACION DE SERVICIOS DE LIMPIEZA\nEntre Oficinas Ejemplo S.L. y Limpiezas Brillo S.L.",
        "Clausula tercera: el proveedor de limpieza realizara la limpieza diaria de las oficinas.\n"
        "Precio mensual: 850 euros mas IVA.",
    ])
    make_pdf(root / "Contabilidad" / "Facturas" / "factura_2024_03_perez.pdf", [
        "FACTURA N 2024-031\nFecha: 15 de marzo de 2024\nCliente: Construcciones Perez\nTotal: 1.250,00 euros",
    ])
    d = docx.Document()
    d.core_properties.title = "Presupuesto reforma"
    d.add_heading("Presupuesto de la reforma de la oficina", 1)
    d.add_paragraph("Demolición de tabiques y retirada de escombros.")
    d.add_paragraph("Alicatado del baño y pintura de paredes.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Partida"
    t.cell(0, 1).text = "Importe"
    t.cell(1, 0).text = "Fontanería"
    t.cell(1, 1).text = "3.200 €"
    d.save(root / "Obras" / "Reforma oficina" / "presupuesto_reforma.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gastos"
    ws.append(["Concepto", "Proveedor", "Importe"])
    ws.append(["Material de oficina", "Papelería Sol", 120])
    ws.append(["Limpieza", "Limpiezas Brillo", 850])
    wb.save(root / "Contabilidad" / "gastos_2024.xlsx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Plan comercial"
    s.placeholders[1].text = "Captar clientes del sector hotelero en la costa"
    prs.save(root / "plan_comercial.pptx")

    m = EmailMessage()
    m["Subject"] = "Renovación del seguro de la furgoneta"
    m["From"] = "Ana <ana@ejemplo.es>"
    m["To"] = "oficina@ejemplo.es"
    m["Date"] = "Mon, 04 Mar 2024 10:00:00 +0100"
    m.set_content("Hola,\n\nOs adjunto la póliza renovada del seguro de la furgoneta.\n\nSaludos")
    m.add_attachment(b"Poliza numero 778899 de la aseguradora Mapfre, vehiculo furgoneta", maintype="text",
                     subtype="plain", filename="poliza.txt")
    (root / "Correo" / "seguro.eml").write_bytes(bytes(m))

    (root / "notas.txt").write_text("Llamar al técnico de la caldera el lunes.\n\nRevisar el extintor.",
                                    encoding="utf-8")
    (root / "clientes.csv").write_text("nombre;ciudad\nHotel Mar Azul;Benidorm\nCamping Pinar;Denia\n",
                                       encoding="cp1252")

    doc = ezdxf.new()
    msp = doc.modelspace()
    doc.layers.add("FONTANERIA")
    msp.add_text("PLANTA BAJA - COCINA", dxfattribs={"layer": "FONTANERIA"})
    msp.add_mtext("Tubería de cobre 22 mm")
    doc.saveas(root / "Obras" / "Reforma oficina" / "plano_planta_baja.dxf")
    (root / "Obras" / "Reforma oficina" / "alzado_fachada.dwg").write_bytes(b"AC1032" + b"\0" * 200)

    make_pdf(root / "RRHH" / "Nominas" / "nomina_marzo_juan.pdf", ["NOMINA marzo Juan Garcia salario 2.100 euros"])
    (root / "Thumbs.db").write_bytes(b"x")
    (root / "~$borrador.docx").write_bytes(b"x")
    (root / "roto.pdf").write_bytes(b"%PDF-1.4 esto no es un pdf")

    if HAS_TESSERACT:
        text_image("FACTURA ESCANEADA\nTaller Mecanico Lopez\nCambio de aceite").save(
            root / "Contabilidad" / "Facturas" / "ticket_taller.png")
        img = text_image("ACTA DE REUNION\nJunta de vecinos del local", size=(1700, 2200))
        img.convert("RGB").save(root / "acta_escaneada.pdf", "PDF", resolution=200)


@pytest.fixture
def corpus(tmp_path) -> Path:
    root = tmp_path / "share"
    root.mkdir()
    build_corpus(root)
    return root


def make_config(tmp_path: Path, root: Path, **perm) -> Config:
    cfg = Config()
    cfg.share.path = str(root)
    cfg.share.unc = r"\\SERVIDOR\Compartido"
    cfg.data_dir = str(tmp_path / "data")
    cfg.model.repo = "hash"
    cfg.model.dim = 256
    cfg.index.extract_workers = 1
    cfg.index.watch = False
    cfg.permissions.mode = perm.get("mode", "ntfs")
    return cfg


@pytest.fixture
def cfg(tmp_path, corpus) -> Config:
    return make_config(tmp_path, corpus)


